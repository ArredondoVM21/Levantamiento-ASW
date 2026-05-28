from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.core.files.base import ContentFile
from django.utils import timezone

# ReportLab (PDF)
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import simpleSplit, ImageReader

# Excel
import openpyxl

# Modelos
from .models import Reporte, CuentaPorPagar
from apphela.models import Presupuesto

# Formularios
from .forms import CuentaPorPagarForm


# INFORMES (LISTADO + FILTROS)
@login_required
def informes(request):
    presupuestos = Presupuesto.objects.all().order_by('-anio', '-mes')

    anio = request.GET.get("anio")
    mes = request.GET.get("mes")

    if anio:
        presupuestos = presupuestos.filter(anio=anio)
    if mes:
        presupuestos = presupuestos.filter(mes=mes)

    anios = Presupuesto.objects.values_list("anio", flat=True).distinct().order_by("anio")
    meses = Presupuesto.objects.values_list("mes", flat=True).distinct().order_by("mes")

    return render(request, 'informes/generar_reportes.html', {
        'presupuestos': presupuestos,
        'anios': anios,
        'meses': meses
    })




# GENERAR PDF CORPORATIVO
@login_required
def generar_pdf(request, id):
    presupuesto = get_object_or_404(Presupuesto, id=id)
    detalles = presupuesto.detalles.all()

    response = HttpResponse(content_type='application/pdf')
    nombre_archivo = f"Reporte_{presupuesto.mes}_{presupuesto.anio}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'

    pdf = canvas.Canvas(response, pagesize=letter)

    # MÁRGENES
    x_left = 50
    x_right = 550
    y = 750

# LOGO + ENCABEZADO
    try:
        logo = ImageReader("https://i.ibb.co/VWRKv90P/SHlogo.png")

        # Tamaño similar al navbar (80x80 px)
        logo_width = 70
        logo_height = 70

        pdf.drawImage(
            logo,
            x_left,             # margen izquierdo
            y - logo_height,    # posición vertical
            width=logo_width,
            height=logo_height,
            mask='auto'
        )
    except:
        pass

    # TÍTULO CORPORATIVO
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(x_left + 90, y - 10, "Secreto - Gestión de Presupuestos")

    # FECHA DE GENERACIÓN
    pdf.setFont("Helvetica", 10)
    pdf.drawString(x_left + 90, y - 25, f"Generado el: {timezone.now().strftime('%d/%m/%Y')}")

    # Línea separadora
    y -= 80
    pdf.setLineWidth(1)
    pdf.line(x_left, y, x_right, y)


# TÍTULO
    y -= 40
    pdf.setFont("Helvetica-Bold", 22)
    pdf.drawCentredString(300, y, f"Reporte de {presupuesto.mes} {presupuesto.anio}")

    y -= 20
    pdf.line(150, y, 450, y)


# DESCRIPCIÓN MULTILÍNEA
    y -= 40
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(x_left, y, "Descripción:")

    descripcion = presupuesto.descripcion
    lineas = simpleSplit(descripcion, "Helvetica", 12, 480)

    y -= 15
    fondo_altura = (len(lineas) * 16) + 10

    # Fondo gris
    pdf.setFillColorRGB(0.95, 0.95, 0.95)
    pdf.rect(x_left - 5, y - fondo_altura + 5, 530, fondo_altura, fill=True, stroke=False)
    pdf.setFillColorRGB(0, 0, 0)

    pdf.setFont("Helvetica", 12)
    for linea in lineas:
        pdf.drawString(x_left, y, linea)
        y -= 16

    y -= 30

 # ENCABEZADO TABLA
    pdf.setFont("Helvetica-Bold", 13)
    pdf.setFillColorRGB(0.05, 0.30, 0.60)  # azul corporativo
    pdf.rect(x_left, y - 5, 500, 22, fill=True, stroke=False)

    pdf.setFillColorRGB(1, 1, 1)
    pdf.drawString(x_left + 10, y, "Items Asociados")
    pdf.drawString(x_left + 350, y, "Monto")
    pdf.setFillColorRGB(0, 0, 0)

    y -= 35


# LISTA ÍTEMS
    pdf.setFont("Helvetica", 12)
    monto_total = 0

    for d in detalles:
        monto_total += d.monto_limite

        # Alternar fondo
        pdf.setFillColorRGB(0.97, 0.97, 0.97) if (d.id % 2 == 0) else pdf.setFillColorRGB(1, 1, 1)
        pdf.rect(x_left, y - 5, 500, 22, fill=True, stroke=False)

        # Texto
        pdf.setFillColorRGB(0, 0, 0)
        pdf.drawString(x_left + 10, y, d.item.nombre)
        pdf.drawRightString(x_left + 480, y, f"${d.monto_limite:,.0f}")

        y -= 28

        if y < 100:
            pdf.showPage()
            y = 750

#  MONTO TOTAL
    pdf.setFont("Helvetica-Bold", 14)
    pdf.setFillColorRGB(0.10, 0.45, 0.80)
    pdf.rect(x_left, y - 5, 500, 25, fill=True, stroke=True)

    pdf.setFillColorRGB(1, 1, 1)
    pdf.drawString(x_left + 10, y, "Monto Total:")
    pdf.drawRightString(x_left + 480, y, f"${monto_total:,.0f}")

    pdf.save()

    # Guardar en BD
    archivo_pdf = response.getvalue()
    reporte = Reporte(
        nombre=f"Presupuesto {presupuesto.mes}-{presupuesto.anio}",
        tipo="PDF",
        generado_por=request.user
    )
    reporte.archivo.save(nombre_archivo, ContentFile(archivo_pdf))

    return response

# GENERAR EXCEL
@login_required
def generar_excel(request, id):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule

    presupuesto = get_object_or_404(Presupuesto, id=id)
    detalles = presupuesto.detalles.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Reporte {presupuesto.mes} {presupuesto.anio}"


# TÍTULO
    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = f"Reporte de {presupuesto.mes} {presupuesto.anio}"
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center")


# ENCABEZADO
    header_fill = PatternFill(start_color="1D5FA7", end_color="1D5FA7", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws["A3"] = "Ítem"
    ws["B3"] = "Monto"

    for col in ["A3", "B3"]:
        ws[col].fill = header_fill
        ws[col].font = header_font
        ws[col].alignment = Alignment(horizontal="center")


# CUERPO DE LA TABLA
    row = 4
    monto_total = 0

    for d in detalles:
        ws[f"A{row}"] = d.item.nombre
        ws[f"B{row}"] = d.monto_limite

        ws[f"B{row}"].number_format = '"$"#,##0'  # formato moneda

        row += 1
        monto_total += d.monto_limite

# FILA MONTO TOTAL
    ws[f"A{row}"] = "Monto Total:"
    ws[f"A{row}"].font = Font(bold=True)

    ws[f"B{row}"] = monto_total
    ws[f"B{row}"].font = Font(bold=True)
    ws[f"B{row}"].number_format = '"$"#,##0'


# FORMATO CONDICIONAL (Montos altos en rojo)
    max_row = row - 1  # última fila de montos

    rojo = Font(color="FF0000", bold=True)

    ws.conditional_formatting.add(
        f"B4:B{max_row}",
        CellIsRule(
            operator='greaterThan',
            formula=['500000'],
            font=rojo
        )
    )

# AUTO-AJUSTAR ANCHO DE COLUMNA
    for col in range(1, 3):
        max_length = 0
        col_letter = get_column_letter(col)

        for cell in ws[col_letter]:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except:
                pass

        ws.column_dimensions[col_letter].width = max_length + 4

# BORDES PARA TODA LA TABLA
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row_cells in ws.iter_rows(min_row=3, max_row=row, min_col=1, max_col=2):
        for cell in row_cells:
            cell.border = thin_border


# RESPUESTA HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    nombre_archivo = f"Reporte_{presupuesto.mes}_{presupuesto.anio}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'

    wb.save(response)

    # Guardar en BD
    reporte = Reporte(
        nombre=f"Presupuesto {presupuesto.mes}-{presupuesto.anio}",
        tipo="Excel",
        generado_por=request.user
    )
    reporte.archivo.save(nombre_archivo, ContentFile(response.getvalue()))

    return response

# Cuentas por Pagar
@login_required
def lista_cuentas(request):

    filtro = request.GET.get('estado', 'todos')

    # Filtrado por estado
    if filtro == 'pendiente':
        cuentas = CuentaPorPagar.objects.filter(estado='pendiente')
    elif filtro == 'pagado':
        cuentas = CuentaPorPagar.objects.filter(estado='pagado')
    elif filtro == 'vencido':
        cuentas = CuentaPorPagar.objects.filter(estado='vencido')
    else:
        cuentas = CuentaPorPagar.objects.all()

    # Orden por fecha de vencimiento
    cuentas = cuentas.order_by('fecha_vencimiento')

    # Acualiza el estado si no esta pagado
    for c in cuentas:
        estado_anterior = c.estado

        nuevo_estado = c.actualizar_estado()

        if estado_anterior != nuevo_estado and estado_anterior != "pagado":
            c.save()

    return render(request, 'cuentas/lista_cuentas.html', {
        'cuentas': cuentas,
        'filtro': filtro
    })

@login_required
def crear_cuenta(request):
    if request.method == 'POST':
        form = CuentaPorPagarForm(request.POST)
        if form.is_valid():
            cuenta = form.save(commit=False)
            cuenta.generado_por = request.user
            cuenta.save()
            messages.success(request, "Cuenta registrada correctamente.")
            return redirect('lista_cuentas')
    else:
        form = CuentaPorPagarForm()

    return render(request, 'cuentas/crear_cuenta.html', {'form': form})


@login_required
def detalle_cuenta(request, pk):
    cuenta = get_object_or_404(CuentaPorPagar, pk=pk)
    return render(request, 'cuentas/detalle_cuenta.html', {'cuenta': cuenta})


@login_required
def editar_cuenta(request, pk):
    cuenta = get_object_or_404(CuentaPorPagar, pk=pk)
    if request.method == 'POST':
        form = CuentaPorPagarForm(request.POST, instance=cuenta)
        if form.is_valid():
            form.save()
            messages.success(request, "Cuenta actualizada.")
            return redirect('lista_cuentas')
    else:
        form = CuentaPorPagarForm(instance=cuenta)

    return render(request, 'cuentas/editar_cuenta.html', {'form': form, 'cuenta': cuenta})


@login_required
def eliminar_cuenta(request, pk):
    cuenta = get_object_or_404(CuentaPorPagar, pk=pk)
    if request.method == 'POST':
        cuenta.delete()
        messages.success(request, "Cuenta eliminada correctamente.")
        return redirect('lista_cuentas')

    return render(request, 'cuentas/eliminar_cuenta.html', {'cuenta': cuenta})
